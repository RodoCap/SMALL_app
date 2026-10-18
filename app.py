import io
import os
import socket
import sqlite3
import traceback
from werkzeug.exceptions import HTTPException
from datetime import date
from flask import Flask, request, jsonify, send_from_directory, send_file
from models import Producto, Inventario
import database
import hmac, time
from datetime import date, timedelta
from flask import session, redirect

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "solo-para-local")
app.config["PERMANENT_SESSION_LIFETIME"] = timedelta(days=30)
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
CLAVE = os.environ.get("SMALL_PASSWORD")  # si no existe (en tu PC), no pide login
app.config["MAX_CONTENT_LENGTH"] = 8 * 1024 * 1024  # fotos de hasta 8 MB
app.config["SEND_FILE_MAX_AGE_DEFAULT"] = 0  # el navegador siempre pide la página al día, sin versiones viejas

CARPETA_FOTOS = os.path.join(app.static_folder, "fotos")
os.makedirs(CARPETA_FOTOS, exist_ok=True)

inventario = Inventario()


# ---------- Respaldo automático ----------
ultimo_respaldo = None

def respaldar():
    # Una copia de la base por día (se hace al arrancar y en la primera petición de cada día)
    global ultimo_respaldo
    try:
        database.hacer_respaldo()
        ultimo_respaldo = date.today()
    except (OSError, sqlite3.Error) as problema:
        print("AVISO: no se pudo hacer el respaldo:", problema)

respaldar()  # antes de crear o modificar tablas: guarda cómo estaba la base al abrir

@app.before_request
def respaldo_diario():
    if ultimo_respaldo != date.today():
        respaldar()


database.crear_tabla()
database.crear_tablas_creditos()
inventario.productos = database.obtener_productos()
inventario.ventas = database.obtener_ventas()
inventario.creditos = database.obtener_creditos(inventario.ventas)


# ---------- Ayudas ----------
def ruta_foto(id_producto):
    return os.path.join(CARPETA_FOTOS, f"{id_producto}.jpg")

def url_foto(id_producto):
    ruta = ruta_foto(id_producto)
    if os.path.exists(ruta):
        # ?v=... obliga al navegador a mostrar la foto nueva si la reemplazas
        return f"/static/fotos/{id_producto}.jpg?v={int(os.path.getmtime(ruta))}"
    return None

def producto_a_dict(p):
    return {
        "id": p.id_producto,
        "nombre": p.nombre,
        "categoria": p.categoria,
        "compra": p.compra,
        "venta": p.venta,
        "cantidad_comprada": p.cantidad_comprada,
        "cantidad_vendida": p.cantidad_vendida,
        "stock_actual": p.stock_actual,
        "total_invertido": p.total_invertido,
        "total_vendido": inventario.total_vendido_de(p.id_producto),
        "ganancia": inventario.ganancia_de(p.id_producto),
        "precio_promedio": inventario.precio_promedio_de(p.id_producto),
        "ganancia_unitaria": p.ganancia_unitaria,
        "foto": url_foto(p.id_producto)
    }

def error(mensaje, codigo):
    return jsonify({"mensaje": mensaje}), codigo


PAGINA_LOGIN = """<!DOCTYPE html>
<html lang="es"><head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">
<meta name="theme-color" content="#FFFFFF">
<title>SmallPhone · Entrar</title>
<style>
  :root { --ink:#3A1F1B; --muted:#7A6560; --paper:#FAF5EE; --card:#FFFDF9; --line:#E7DCCF;
          --accent:#C93D30; --brand:#E0352A; --danger:#A12A20; }
  * { box-sizing: border-box; }
  body { margin:0; min-height:100dvh; display:grid; place-items:center; padding:16px;
         font-family: system-ui,-apple-system,"Segoe UI",Roboto,sans-serif;
         background:var(--paper); color:var(--ink); }
  form { width:min(100%,360px); display:grid; gap:14px; background:var(--card);
         border:1px solid var(--line); border-radius:14px; padding:24px; }
  h1 { margin:0; text-align:center; font-size:1.5rem; color:var(--brand); }
  p.sub { margin:-6px 0 0; text-align:center; color:var(--muted); font-size:.9rem; }
  input { font:inherit; font-size:16px; color:var(--ink); background:var(--paper);
          border:1px solid var(--line); border-radius:10px; padding:0 12px; min-height:48px; width:100%; }
  input:focus-visible { outline:3px solid var(--accent); outline-offset:2px; }
  button { font:inherit; font-weight:650; min-height:48px; border-radius:10px; border:0;
           background:var(--accent); color:#fff; cursor:pointer; }
  .error { margin:0; text-align:center; color:var(--danger); font-weight:600; min-height:1.2em; }
</style></head>
<body>
<form method="post">
  <h1>SmallPhone</h1>
  <p class="sub">Ingresa tu clave</p>
  <input type="text" name="usuario" value="smallphone" autocomplete="username" hidden>
  <input type="password" name="clave" placeholder="Clave" autocomplete="current-password" autofocus required>
  <button type="submit">Entrar</button>
  <p class="error">MENSAJE</p>
</form>
</body></html>"""

@app.route("/login", methods=["GET", "POST"])
def login():
    mensaje = ""
    if request.method == "POST":
        if CLAVE and hmac.compare_digest(request.form.get("clave", ""), CLAVE):
            session.permanent = True
            session["ok"] = True
            return redirect("/")
        time.sleep(1)  # frena a quien intente adivinar la clave
        mensaje = "Clave incorrecta"
    return PAGINA_LOGIN.replace("MENSAJE", mensaje)

@app.route("/salir")
def salir():
    session.clear()
    return redirect("/login")

@app.before_request
def exigir_login():
    if not CLAVE or request.endpoint == "login" or session.get("ok"):
        return
    if request.path == "/":
        return redirect("/login")
    return error("No autorizado", 401)


@app.errorhandler(413)
def foto_muy_pesada(e):
    return error("La foto pesa demasiado (máximo 8 MB)", 413)

@app.errorhandler(Exception)
def error_inesperado(e):
    # Cualquier error que no esperábamos: lo imprime en la terminal y lo devuelve
    # como JSON con el nombre del error (así se ve en Thunder Client)
    if isinstance(e, HTTPException):
        return e
    traceback.print_exc()
    return error(f"Error interno: {type(e).__name__}: {e}", 500)


# ---------- Página ----------
@app.route("/")
def inicio():
    return send_from_directory("static", "index.html")


# ---------- Productos ----------
@app.route("/productos", methods=["GET"])
def obtener_productos():
    return jsonify([producto_a_dict(p) for p in inventario.productos])

@app.route("/productos", methods=["POST"])
def crear_producto():
    datos = request.get_json(silent=True) or {}
    try:
        if inventario.buscar_producto(datos["id_producto"]):
            return error("Ya existe un producto con ese ID", 409)
        if inventario.tuvo_ventas(datos["id_producto"]):
            return error("Ese ID perteneció a un producto que ya tuvo ventas. Usa otro ID", 409)
        nuevo = Producto(
            nombre=datos["nombre"],
            compra=datos["compra"],
            venta=datos["venta"],
            cantidad_comprada=datos["cantidad_comprada"],
            id_producto=datos["id_producto"],
            categoria=datos["categoria"]
        )
    except KeyError as falta:
        return error(f"Falta el dato {falta}", 400)
    except ValueError as problema:
        return error(str(problema), 400)
    except TypeError:
        return error("Hay un dato con formato inválido", 400)
    inventario.agregar_producto(nuevo)
    database.insertar_producto(nuevo)
    return jsonify({"mensaje": "Producto creado"}), 201

@app.route("/productos/<int:id_producto>", methods=["PUT"])
def editar_producto(id_producto):
    producto = inventario.buscar_producto(id_producto)
    if producto is None:
        return error("Producto no encontrado", 404)
    datos = request.get_json(silent=True) or {}
    try:
        producto.actualizar(
            nombre=datos["nombre"],
            categoria=datos["categoria"],
            compra=datos["compra"],
            venta=datos["venta"],
            cantidad_comprada=datos.get("cantidad_comprada")  # opcional: corrige el total comprado
        )
    except KeyError as falta:
        return error(f"Falta el dato {falta}", 400)
    except ValueError as problema:
        return error(str(problema), 400)
    except TypeError:
        return error("Hay un dato con formato inválido", 400)
    database.guardar_producto(producto)
    return jsonify({"mensaje": "Producto actualizado"}), 200

@app.route("/productos/<int:id_producto>/reponer", methods=["POST"])
def reponer_producto(id_producto):
    producto = inventario.buscar_producto(id_producto)
    if producto is None:
        return error("Producto no encontrado", 404)
    datos = request.get_json(silent=True) or {}
    try:
        producto.reponer(datos["cantidad"], datos.get("costo"))
    except KeyError:
        return error("Falta la cantidad", 400)
    except ValueError as problema:
        return error(str(problema), 400)
    database.guardar_producto(producto)
    return jsonify({"mensaje": "Stock repuesto", "stock_actual": producto.stock_actual}), 200

@app.route("/productos/<int:id_producto>", methods=["DELETE"])
def eliminar_producto(id_producto):
    eliminado = inventario.eliminar_producto(id_producto)
    if eliminado:
        database.eliminar_producto(id_producto)
        # Borra también su foto: si luego se reutiliza el ID, no hereda la foto vieja
        if os.path.exists(ruta_foto(id_producto)):
            os.remove(ruta_foto(id_producto))
        return jsonify({"mensaje": "Producto eliminado"}), 200
    return error("Producto no encontrado", 404)


# ---------- Foto del producto ----------
@app.route("/productos/<int:id_producto>/foto", methods=["POST"])
def subir_foto(id_producto):
    if inventario.buscar_producto(id_producto) is None:
        return error("Producto no encontrado", 404)
    archivo = request.files.get("foto")
    if archivo is None:
        return error("No llegó ninguna foto", 400)
    contenido = archivo.read()
    if not contenido.startswith(b"\xff\xd8\xff"):  # firma de un archivo JPG
        return error("La foto debe ser un JPG", 400)
    with open(ruta_foto(id_producto), "wb") as f:
        f.write(contenido)
    return jsonify({"mensaje": "Foto guardada", "foto": url_foto(id_producto)}), 200


# ---------- Ventas ----------
@app.route("/ventas", methods=["GET"])
def obtener_ventas():
    return jsonify([
        {
            "id_venta": v.id_venta,
            "fecha": v.fecha,
            "id_producto": v.id_producto,
            "producto": v.producto,
            "cantidad": v.cantidad_vendida,
            "precio_unitario": v.precio_unitario,
            "total": v.total,
            "ganancia": v.ganancia,
            "costo_unitario": v.costo_unitario,
            "id_credito": v.id_credito
        }
        for v in inventario.ventas
    ])

@app.route("/ventas", methods=["POST"])
def crear_venta():
    datos = request.get_json(silent=True) or {}
    try:
        producto = inventario.buscar_producto(datos["id_producto"])
        resultado = inventario.registrar_venta(
            id_producto=datos["id_producto"],
            cantidad=datos["cantidad"],
            precio=datos.get("precio")  # opcional: si no llega, usa el precio de lista
        )
    except KeyError as falta:
        return error(f"Falta el dato {falta}", 400)
    except ValueError as problema:
        return error(str(problema), 400)
    except TypeError:
        return error("Hay un dato con formato inválido", 400)
    if not resultado:
        return error("Producto no encontrado", 404)

    ultima_venta = inventario.ventas[-1]
    try:
        database.guardar_venta(ultima_venta, producto)
    except sqlite3.Error:
        # Si la base de datos falla, deshace el cambio en memoria para que no queden desfasados
        producto.cantidad_vendida -= ultima_venta.cantidad_vendida
        inventario.ventas.pop()
        return error("No se pudo guardar la venta. Intenta de nuevo", 500)
    return jsonify({"mensaje": "Venta registrada"}), 201


@app.route("/ventas/<int:id_venta>", methods=["DELETE"])
def anular_venta(id_venta):
    venta = inventario.buscar_venta(id_venta)
    if venta is None:
        return error("Venta no encontrada", 404)
    if venta.id_credito is not None:
        return error("Esta venta pertenece a un crédito. Anula el crédito completo desde Créditos", 409)
    try:
        database.anular_venta(venta.id_venta, venta.id_producto, venta.cantidad_vendida)
    except sqlite3.Error:
        return error("No se pudo anular la venta. Intenta de nuevo", 500)
    inventario.anular_venta(id_venta)  # recién ahora, con la base de datos ya actualizada
    return jsonify({"mensaje": "Venta anulada"}), 200


# ---------- Créditos ----------
@app.route("/creditos", methods=["GET"])
def obtener_creditos():
    return jsonify([
        {
            "id_credito": c.id_credito,
            "cliente": c.cliente,
            "fecha": c.fecha,
            "fecha_proximo_pago": c.fecha_proximo_pago,
            "total": c.total,
            "total_abonado": c.total_abonado,
            "saldo": c.saldo,
            "esta_pagado": c.esta_pagado,
            "vencido": c.vencido,
            "productos": [
                {"producto": v.producto, "cantidad": v.cantidad_vendida,
                 "precio_unitario": v.precio_unitario, "total": v.total}
                for v in c.ventas
            ],
            "abonos": [
                {"id_abono": a.id_abono, "fecha": a.fecha, "monto": a.monto}
                for a in c.abonos
            ]
        }
        for c in inventario.creditos
    ])

@app.route("/creditos", methods=["POST"])
def crear_credito():
    datos = request.get_json(silent=True) or {}
    try:
        credito = inventario.registrar_credito(
            cliente=datos.get("cliente"),
            items=datos.get("items"),
            abono_inicial=datos.get("abono") or 0,
            fecha_proximo_pago=datos.get("fecha_proximo_pago") or None
        )
    except ValueError as problema:
        return error(str(problema), 400)
    try:
        database.guardar_credito(credito)
    except sqlite3.Error:
        # Si la base falla, deshace los cambios en memoria
        for v in credito.ventas:
            inventario.buscar_producto(v.id_producto).cantidad_vendida -= v.cantidad_vendida
            inventario.ventas.remove(v)
        inventario.creditos.remove(credito)
        return error("No se pudo guardar el crédito. Intenta de nuevo", 500)
    return jsonify({"mensaje": "Crédito registrado", "id_credito": credito.id_credito}), 201

@app.route("/creditos/<int:id_credito>/abonos", methods=["POST"])
def crear_abono(id_credito):
    datos = request.get_json(silent=True) or {}
    try:
        abono = inventario.registrar_abono(id_credito, datos.get("monto"))
    except ValueError as problema:
        return error(str(problema), 400)
    if abono is None:
        return error("Crédito no encontrado", 404)
    credito = inventario.buscar_credito(id_credito)
    try:
        database.guardar_abono(abono)
    except sqlite3.Error:
        credito.abonos.remove(abono)
        return error("No se pudo guardar el abono. Intenta de nuevo", 500)
    return jsonify({"mensaje": "Abono registrado", "saldo": credito.saldo}), 201

@app.route("/creditos/<int:id_credito>", methods=["DELETE"])
def anular_credito(id_credito):
    credito = inventario.buscar_credito(id_credito)
    if credito is None:
        return error("Crédito no encontrado", 404)
    try:
        database.anular_credito(credito)
    except sqlite3.Error:
        return error("No se pudo anular el crédito. Intenta de nuevo", 500)
    inventario.anular_credito(id_credito)  # recién ahora, con la base ya actualizada
    return jsonify({"mensaje": "Crédito anulado"}), 200


# ---------- Exportar a Excel ----------
@app.route("/exportar", methods=["GET"])
def exportar_excel():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        return error("Falta instalar openpyxl. En la terminal escribe: pip install openpyxl", 500)

    def llenar(hoja, titulos, filas, formatos):
        hoja.append(titulos)
        for celda in hoja[1]:
            celda.font = Font(bold=True)
        for fila in filas:
            hoja.append(fila)
        for columna, formato in enumerate(formatos, start=1):
            if formato:
                for celda in hoja.iter_cols(min_col=columna, max_col=columna, min_row=2):
                    for c in celda:
                        c.number_format = formato
        for columna in hoja.columns:
            ancho = max(len(str(c.value)) if c.value is not None else 0 for c in columna)
            hoja.column_dimensions[columna[0].column_letter].width = min(max(ancho + 3, 11), 40)
        hoja.freeze_panes = "A2"

    def como_fecha(texto):
        try:
            return date.fromisoformat(texto)
        except (TypeError, ValueError):
            return texto

    libro = Workbook()
    hoja_ventas = libro.active
    hoja_ventas.title = "Ventas"
    llenar(hoja_ventas,
           ["ID venta", "Fecha", "Producto", "Cantidad", "Precio unitario", "Total", "Costo unitario", "Ganancia"],
           [[v.id_venta, como_fecha(v.fecha), v.producto, v.cantidad_vendida, v.precio_unitario,
             v.total, v.costo_unitario, v.ganancia] for v in inventario.ventas],
           [None, "yyyy-mm-dd", None, "#,##0", "#,##0", "#,##0", "#,##0", "#,##0"])

    hoja_productos = libro.create_sheet("Productos")
    llenar(hoja_productos,
           ["ID", "Nombre", "Categoría", "Costo", "Precio de lista", "Comprado", "Vendido", "Stock",
            "Total invertido", "Total vendido", "Ganancia", "Precio promedio vendido"],
           [[p.id_producto, p.nombre, p.categoria, p.compra, p.venta, p.cantidad_comprada,
             p.cantidad_vendida, p.stock_actual, p.total_invertido,
             inventario.total_vendido_de(p.id_producto), inventario.ganancia_de(p.id_producto),
             inventario.precio_promedio_de(p.id_producto)] for p in inventario.productos],
           [None, None, None, "#,##0", "#,##0", "#,##0", "#,##0", "#,##0", "#,##0", "#,##0", "#,##0", "#,##0"])

    archivo = io.BytesIO()
    libro.save(archivo)
    archivo.seek(0)
    return send_file(
        archivo, as_attachment=True,
        download_name=f"SmallPhone_{date.today().isoformat()}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ---------- Resumen ----------
@app.route("/resumen", methods=["GET"])
def obtener_resumen():
    total_invertido = sum(p.total_invertido for p in inventario.productos)
    total_vendido = sum(v.total for v in inventario.ventas)
    # La ganancia sale de las ventas (guardan el precio y costo de ese día)
    ganancia_total = sum(v.ganancia for v in inventario.ventas)

    return jsonify({
        "total_invertido": total_invertido,
        "total_vendido": total_vendido,
        "ganancia_total": ganancia_total
    })


def ip_local():
    # La dirección de tu PC en el WiFi, para abrir la página desde el celular
    s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    try:
        s.connect(("10.255.255.255", 1))
        return s.getsockname()[0]
    except OSError:
        return "127.0.0.1"
    finally:
        s.close()


if __name__ == "__main__":
    # debug=False: con host 0.0.0.0 el modo debug quedaría abierto a toda tu red WiFi.
    # Si vas a programar cambios, ponlo en True solo mientras desarrollas.
    print(f"\nEn el PC:      http://127.0.0.1:5000")
    print(f"En el celular: http://{ip_local()}:5000\n")
    app.run(host="0.0.0.0", port=5000, debug=False)